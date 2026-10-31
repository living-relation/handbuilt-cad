"""Tests for the pure (no adsk) parts of the Fusion add-ins in tools/fusion_addins."""
import csv
import io
import json
import os
import sys
import zipfile

import pytest

ROOT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'tools', 'fusion_addins')
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, 'HBExport'))
sys.path.insert(0, os.path.join(ROOT, 'HBBom'))

import bom_core  # noqa: E402
import export_core  # noqa: E402
from hbcad_fusion_common import settings  # noqa: E402
from hbcad_fusion_common.names import safe_name  # noqa: E402


def test_safe_name_strips_version_and_bad_chars():
    assert safe_name('Filament Cabinet v12') == 'Filament Cabinet'
    assert safe_name('a/b:c*?') == 'a_b_c__'
    assert safe_name('  ') == 'export'


def test_target_paths_order_and_overwrite_names(tmp_path):
    p = export_core.target_paths(str(tmp_path), 'Part v3', ['dxf', 'step', 'stl'])
    assert list(p) == ['step', 'stl', 'dxf']
    assert p['step'].endswith('Part.step')
    assert export_core.target_paths(str(tmp_path), 'Part v4', ['step']) == {'step': p['step']}


def test_summarize_marks_failures():
    txt = export_core.summarize({'iges': (False, 'blocked'), 'step': (True, 'x.step')})
    assert txt.splitlines() == ['OK      STEP: x.step', 'FAILED  IGES: blocked']


def test_settings_roundtrip_ignores_unknown_and_bad(tmp_path):
    settings.save('export', {'folder': 'C:/x', 'junk': 1}, str(tmp_path))
    got = settings.load('export', export_core.DEFAULTS, str(tmp_path))
    assert got['folder'] == 'C:/x' and 'junk' not in got
    (tmp_path / 'bad.json').write_text('{oops')
    assert settings.load('bad', {'a': 1}, str(tmp_path)) == {'a': 1}


def test_count_parts():
    assert bom_core.count_parts(['a', 'b', 'a']) == {'a': 2, 'b': 1}


@pytest.mark.parametrize('data', [
    {'parts': [{'Name': 'Leg', 'PN': 'FC5-1', 'Material': 'Aluminum 6061', 'McMaster': '91290A'}]},
    [{'component': 'Leg', 'part_number': 'FC5-1', 'material': 'Aluminum 6061', 'vendor_pn': '91290A'}],
    {'Leg': {'pn': 'FC5-1', 'mat': 'Aluminum 6061', 'mcmaster_pn': '91290A'}},
])
def test_parse_build_json_shapes(data):
    t = bom_core.parse_build_json(data)
    assert t['leg']['part_number'] == 'FC5-1'
    assert t['fc5-1']['material'] == 'Aluminum 6061'
    assert t['leg']['vendor_pn'] == '91290A'


def test_fill_from_build_respects_overwrite():
    rows = [{'name': 'Leg', 'qty': 4, 'part_number': 'KEEP', 'description': ''},
            {'name': 'Other', 'qty': 1}]
    t = bom_core.parse_build_json({'Leg': {'pn': 'NEW', 'description': 'leg'}})
    got = bom_core.fill_from_build(rows, t)
    assert rows[0]['part_number'] == 'KEEP' and rows[0]['description'] == 'leg'
    assert got == [('Leg', 'description')]
    bom_core.fill_from_build(rows, t, overwrite=True)
    assert rows[0]['part_number'] == 'NEW'


def test_fill_replaces_fusion_defaults():
    rows = [{'name': 'Leg', 'part_number': 'Leg', 'material': 'Steel', '_soft': ['part_number', 'material']}]
    t = bom_core.parse_build_json({'Leg': {'pn': 'FC5-1', 'material': 'Aluminum 6061'}})
    assert sorted(f for _, f in bom_core.fill_from_build(rows, t)) == ['material', 'part_number']
    assert rows[0]['part_number'] == 'FC5-1' and rows[0]['_soft'] == []


def test_empty_and_report():
    rows = [{'name': 'Leg', 'qty': 2, 'part_number': 'P', 'description': '', 'material': 'Steel'}]
    assert bom_core.empty_cells(rows) == [('Leg', 'description')]
    rep = bom_core.report(rows, [('Leg', 'material', 'no material')], [])
    assert '2 pieces' in rep and 'Leg: material - no material' in rep and 'Leg: description' in rep


def test_csv_and_xlsx(tmp_path):
    rows = [{'name': 'A & <B>', 'qty': 3, 'part_number': 'P1'}]
    out = list(csv.reader(io.StringIO(bom_core.to_csv(rows, ['name', 'qty', 'part_number']))))
    assert out[0] == ['QTY', 'Part Number', 'Component'] and out[1] == ['3', 'P1', 'A & <B>']
    p = bom_core.write_xlsx(str(tmp_path / 'b.xlsx'), rows, columns=['qty', 'name'])
    with zipfile.ZipFile(p) as z:
        sheet = z.read('xl/worksheets/sheet1.xml').decode()
        assert 'A &amp; &lt;B&gt;' in sheet and '<v>3</v>' in sheet
        assert '[Content_Types].xml' in z.namelist()
    try:
        import openpyxl
    except ImportError:
        return
    ws = openpyxl.load_workbook(p).active
    assert ws['A2'].value == 3 and ws['B2'].value == 'A & <B>'


def test_col_letters():
    assert [bom_core._col(i) for i in (0, 25, 26, 27)] == ['A', 'Z', 'AA', 'AB']


def test_manifests_valid():
    for name in ('HBExport', 'HBBom'):
        with open(os.path.join(ROOT, name, name + '.manifest'), encoding='utf-8') as f:
            m = json.load(f)
        assert m['type'] == 'addin' and m['autodeskProduct'] == 'Fusion'


def test_icons_present():
    import pathlib
    from PIL import Image
    root = pathlib.Path(__file__).resolve().parents[1] / 'tools' / 'fusion_addins'
    for a in ('HBExport', 'HBBom'):
        d = root / a / 'resources' / 'icon'
        for name, n in (('16x16', 16), ('32x32', 32)):
            assert Image.open(d / f'{name}.png').size == (n, n)
        assert "'resources', 'icon'" in (root / a / f'{a}.py').read_text()


def _n(i, kids=(), pn=''):
    return {'id': i, 'name': i, 'part_number': pn or i.upper(), 'description': 'd' + i, 'children': list(kids)}


# root: 2x SubA (each: 4x screw, 1x plate), 1x Base, 3x screw
TREE = [_n('suba', [_n('screw'), _n('screw'), _n('screw'), _n('screw'), _n('plate')]),
        _n('base'),
        _n('suba', [_n('screw'), _n('screw'), _n('screw'), _n('screw'), _n('plate')]),
        _n('screw'), _n('screw'), _n('screw')]


def test_column_order_and_no_uom():
    assert bom_core.KEYS[:4] == ['item', 'qty', 'part_number', 'description']
    assert not any('unit' in k or 'uom' in k for k in bom_core.KEYS)
    assert 'item' not in bom_core.EDITABLE and 'qty' not in bom_core.EDITABLE
    assert bom_core.pick_columns(['notes', 'qty', 'item', 'bogus']) == ['item', 'qty', 'notes']
    assert bom_core.pick_columns([]) == bom_core.DEFAULT_COLUMNS


def test_top_level_style():
    rows = bom_core.build_rows(TREE, 'top')
    assert [(r['item'], r['id'], r['qty']) for r in rows] == [('1', 'suba', 2), ('2', 'base', 1), ('3', 'screw', 3)]
    assert all(r['top'] for r in rows) and rows[0]['has_children']


def test_indented_style():
    rows = bom_core.build_rows(TREE, 'indented')
    assert [(r['item'], r['id'], r['qty'], r['level']) for r in rows] == [
        ('1', 'suba', 2, 0), ('1.1', 'screw', 4, 1), ('1.2', 'plate', 1, 1),
        ('2', 'base', 1, 0), ('3', 'screw', 3, 0)]
    assert [r['id'] for r in bom_core.thumb_rows(rows)] == ['suba', 'base', 'screw']


def test_parts_only_style_totals():
    rows = bom_core.build_rows(TREE, 'parts')
    assert [(r['item'], r['id'], r['qty']) for r in rows] == [('1', 'screw', 11), ('2', 'plate', 2), ('3', 'base', 1)]


def test_edits_merge_across_duplicate_rows():
    rows = bom_core.build_rows(TREE, 'indented')
    rows[4]['part_number'] = '91290A115'  # top-level screw row edited
    merged = bom_core.merge_edits(rows)
    assert len(merged) == 4
    assert next(m for m in merged if m['id'] == 'screw')['part_number'] == '91290A115'
    new = bom_core.spread_edits(bom_core.build_rows(TREE, 'top'), merged)
    assert new[2]['part_number'] == '91290A115' and new[0]['part_number'] == 'SUBA'


def test_export_follows_style_and_columns(tmp_path):
    rows = bom_core.build_rows(TREE, 'indented')
    out = list(csv.reader(io.StringIO(bom_core.to_csv(rows, ['description', 'item', 'qty', 'part_number']))))
    assert out[0] == ['Item', 'QTY', 'Part Number', 'Part Description']
    assert out[2] == ['1.1', '4', 'SCREW', 'dscrew']
    p = bom_core.write_xlsx(str(tmp_path / 'i.xlsx'), rows, columns=['item', 'qty'])
    with zipfile.ZipFile(p) as z:
        sheet = z.read('xl/worksheets/sheet1.xml').decode()
    assert '>1.1<' in sheet and 'thumb' not in sheet


def test_pin_not_forced_each_start():
    import pathlib
    root = pathlib.Path(__file__).resolve().parents[1] / 'tools' / 'fusion_addins'
    for a in ('HBExport', 'HBBom'):
        src = (root / a / f'{a}.py').read_text()
        assert 'ctl.isPromoted =' not in src and 'isPromotedByDefault' in src
